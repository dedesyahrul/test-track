from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from app.database import get_db
from app.models.models import Defect, Module, SubModule, ReportSummary, DefectScoring, TestScript, SitConfig
from typing import List, Optional
from datetime import date, datetime
from io import BytesIO

# Openpyxl for Excel Report
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Reportlab for PDF Report
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

router = APIRouter()


@router.get("/summary-by-submodule")
def get_summary_by_submodule(db: Session = Depends(get_db)):
    results = db.query(
        SubModule.name.label('sub_module'),
        Module.name.label('module'),
        ReportSummary.total_defect,
        ReportSummary.total_non_defect,
        ReportSummary.total
    ).join(
        ReportSummary, ReportSummary.sub_module_id == SubModule.id
    ).join(
        Module, Module.id == SubModule.module_id
    ).filter(
        ReportSummary.total > 0
    ).order_by(Module.name, SubModule.name).all()

    return [
        {
            "sub_module": r.sub_module,
            "module": r.module,
            "total_defect": r.total_defect,
            "total_non_defect": r.total_non_defect,
            "total": r.total
        }
        for r in results
    ]


@router.get("/scoring-summary")
def get_scoring_summary(db: Session = Depends(get_db)):
    scores = db.query(DefectScoring).order_by(DefectScoring.weight.desc()).all()
    total_closed = sum(s.total_closed for s in scores)
    total_open = sum(s.total_open for s in scores)
    total_score_open = sum(s.score_open for s in scores)

    return {
        "items": [
            {
                "category": s.category,
                "weight": s.weight,
                "total_closed": s.total_closed,
                "total_open": s.total_open,
                "score_open": s.score_open
            }
            for s in scores
        ],
        "totals": {
            "total_closed": total_closed,
            "total_open": total_open,
            "total_score_open": total_score_open,
            "total_all": total_closed + total_open
        }
    }


@router.get("/daily-summary")
def get_daily_summary(db: Session = Depends(get_db)):
    daily = db.query(
        Defect.date_created,
        func.count(Defect.id).label('total'),
        func.sum(
            case(
                (Defect.defect_criteria == 'Defect', 1),
                else_=0
            )
        ).label('defects'),
        func.sum(
            case(
                (Defect.defect_criteria == 'Non-Defect', 1),
                else_=0
            )
        ).label('non_defects')
    ).filter(
        Defect.date_created.isnot(None)
    ).group_by(Defect.date_created).order_by(Defect.date_created).all()

    return [
        {
            "date": str(d.date_created),
            "total": d.total,
            "defects": int(d.defects or 0),
            "non_defects": int(d.non_defects or 0)
        }
        for d in daily
    ]


@router.get("/module-summary")
def get_module_summary(db: Session = Depends(get_db)):
    modules = db.query(Module).all()
    result = []

    for mod in modules:
        subs = db.query(SubModule).filter(SubModule.module_id == mod.id).all()
        sub_details = []
        mod_total_defect = 0
        mod_total_non_defect = 0

        for sub in subs:
            defect_count = db.query(func.count(Defect.id)).filter(
                Defect.sub_module_id == sub.id,
                Defect.defect_criteria == 'Defect'
            ).scalar() or 0
            non_defect_count = db.query(func.count(Defect.id)).filter(
                Defect.sub_module_id == sub.id,
                Defect.defect_criteria == 'Non-Defect'
            ).scalar() or 0

            if defect_count > 0 or non_defect_count > 0:
                sub_details.append({
                    "name": sub.name,
                    "defects": defect_count,
                    "non_defects": non_defect_count,
                    "total": defect_count + non_defect_count
                })
                mod_total_defect += defect_count
                mod_total_non_defect += non_defect_count

        if mod_total_defect > 0 or mod_total_non_defect > 0:
            result.append({
                "module": mod.name,
                "total_defect": mod_total_defect,
                "total_non_defect": mod_total_non_defect,
                "total": mod_total_defect + mod_total_non_defect,
                "sub_modules": sub_details
            })

    return result


@router.get("/pdf")
def generate_pdf_report(
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    db: Session = Depends(get_db)
):
    """Generate Executive SIT Summary PDF Report."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'TitleStyle', parent=styles['Heading1'],
        fontName='Helvetica-Bold', fontSize=18, textColor=colors.HexColor('#1E3A5F'),
        spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'SubTitleStyle', parent=styles['Normal'],
        fontName='Helvetica', fontSize=10, textColor=colors.HexColor('#64748B'),
        spaceAfter=15
    )
    section_style = ParagraphStyle(
        'SectionStyle', parent=styles['Heading2'],
        fontName='Helvetica-Bold', fontSize=12, textColor=colors.HexColor('#1E3A5F'),
        spaceBefore=12, spaceAfter=6
    )
    body_style = ParagraphStyle(
        'BodyStyle', parent=styles['Normal'],
        fontName='Helvetica', fontSize=9, leading=11, textColor=colors.HexColor('#334155')
    )
    table_header_style = ParagraphStyle(
        'TableHeader', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=9, textColor=colors.white, alignment=1
    )
    table_body_style = ParagraphStyle(
        'TableBody', parent=styles['Normal'],
        fontName='Helvetica', fontSize=8, leading=10, textColor=colors.HexColor('#1E293B')
    )

    elements = []

    # Title Banner
    elements.append(Paragraph("EXECUTIVE SIT TESTING REPORT", title_style))
    elements.append(Paragraph(f"Procurement Management System &bull; Exported Date: {date.today().strftime('%d %B %Y')}", subtitle_style))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#1E3A5F'), spaceAfter=12))

    # KPI Statistics Summary
    total_defects = db.query(func.count(Defect.id)).filter(Defect.defect_criteria == 'Defect').scalar() or 0
    open_defects = db.query(func.count(Defect.id)).filter(Defect.defect_criteria == 'Defect', Defect.status.in_(['Open', 'Re-Opened'])).scalar() or 0
    closed_defects = db.query(func.count(Defect.id)).filter(Defect.defect_criteria == 'Defect', Defect.status == 'Closed').scalar() or 0
    resolution_rate = round((closed_defects / total_defects * 100), 1) if total_defects > 0 else 0.0

    # Test Script Stats
    ts_total = db.query(func.count(TestScript.id)).scalar() or 0
    ts_pass = db.query(func.count(TestScript.id)).filter(TestScript.status_by_tester.ilike("pass%")).scalar() or 0
    ts_pass_rate = round((ts_pass / ts_total * 100), 1) if ts_total > 0 else 0.0

    kpi_data = [
        [
            Paragraph(f"<b>Total Defects</b><br/><font size=14 color='#1E3A5F'><b>{total_defects}</b></font>", body_style),
            Paragraph(f"<b>Open Defects</b><br/><font size=14 color='#EF4444'><b>{open_defects}</b></font>", body_style),
            Paragraph(f"<b>Closed Defects</b><br/><font size=14 color='#22C55E'><b>{closed_defects}</b></font>", body_style),
            Paragraph(f"<b>Resolution Rate</b><br/><font size=14 color='#2563EB'><b>{resolution_rate}%</b></font>", body_style),
            Paragraph(f"<b>Test Pass Rate</b><br/><font size=14 color='#10B981'><b>{ts_pass_rate}%</b></font>", body_style),
        ]
    ]
    kpi_table = Table(kpi_data, colWidths=[1.05*inch]*5)
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#E2E8F0')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 8),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 15))

    # Section 1: Scoring Level Summary Table
    elements.append(Paragraph("1. Pembobotan & Penyelesaian Defect", section_style))
    scores = db.query(DefectScoring).order_by(DefectScoring.weight.desc()).all()
    scoring_data = [
        [
            Paragraph("Category", table_header_style),
            Paragraph("Weight", table_header_style),
            Paragraph("Total Closed", table_header_style),
            Paragraph("Total Open", table_header_style),
            Paragraph("Score Open (Open x Weight)", table_header_style)
        ]
    ]
    tot_c, tot_o, tot_s = 0, 0, 0
    for s in scores:
        tot_c += s.total_closed
        tot_o += s.total_open
        tot_s += s.score_open
        scoring_data.append([
            Paragraph(s.category, table_body_style),
            Paragraph(str(s.weight), table_body_style),
            Paragraph(str(s.total_closed), table_body_style),
            Paragraph(str(s.total_open), table_body_style),
            Paragraph(str(s.score_open), table_body_style),
        ])
    scoring_data.append([
        Paragraph("<b>TOTAL</b>", table_body_style),
        Paragraph("-", table_body_style),
        Paragraph(f"<b>{tot_c}</b>", table_body_style),
        Paragraph(f"<b>{tot_o}</b>", table_body_style),
        Paragraph(f"<b>{tot_s}</b>", table_body_style),
    ])

    sc_table = Table(scoring_data, colWidths=[1.3*inch, 0.9*inch, 1.3*inch, 1.3*inch, 1.8*inch])
    sc_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E3A5F')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#F1F5F9')),
        ('ALIGN', (1,1), (-1,-1), 'CENTER'),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
    ]))
    elements.append(sc_table)
    elements.append(Spacer(1, 15))

    # Section 2: Summary Defect per Module
    elements.append(Paragraph("2. Summary Defect per Module", section_style))
    mod_summary = get_module_summary(db)
    mod_data = [
        [
            Paragraph("Module Name", table_header_style),
            Paragraph("Defects", table_header_style),
            Paragraph("Non-Defects", table_header_style),
            Paragraph("Total Items", table_header_style)
        ]
    ]
    for m in mod_summary:
        mod_data.append([
            Paragraph(m['module'], table_body_style),
            Paragraph(str(m['total_defect']), table_body_style),
            Paragraph(str(m['total_non_defect']), table_body_style),
            Paragraph(str(m['total']), table_body_style)
        ])

    mod_table = Table(mod_data, colWidths=[3.6*inch, 1.0*inch, 1.0*inch, 1.0*inch])
    mod_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E3A5F')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('ALIGN', (1,1), (-1,-1), 'CENTER'),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
    ]))
    elements.append(mod_table)

    # Build PDF
    doc.build(elements)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=Executive_SIT_Report_{date.today().strftime('%Y%m%d')}.pdf"}
    )


@router.get("/export-excel")
def export_executive_excel(db: Session = Depends(get_db)):
    """Export Executive SIT Report (Multi-sheet Excel)."""
    wb = Workbook()
    
    # Styles
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    data_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", bold=True, size=10)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    # Sheet 1: Executive Summary
    ws1 = wb.active
    ws1.title = "Executive Summary"

    ws1.cell(row=1, column=1, value="EXECUTIVE SIT REPORT").font = Font(name="Calibri", bold=True, size=16, color="1E3A5F")
    ws1.cell(row=2, column=1, value=f"Project: Procurement Management System | Date: {date.today().strftime('%d-%m-%Y')}").font = Font(size=11, color="64748B")

    # Table 1: Scoring Summary
    ws1.cell(row=4, column=1, value="Pembobotan & Penyelesaian Defect").font = Font(bold=True, size=12)
    headers_s1 = ["Kategori", "Bobot Penilaian", "Jumlah Status Closed", "Jumlah Status Open", "Score Defect Open"]
    for col_idx, h in enumerate(headers_s1, 1):
        cell = ws1.cell(row=5, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    scores = db.query(DefectScoring).order_by(DefectScoring.weight.desc()).all()
    r_idx = 6
    for s in scores:
        vals = [s.category, s.weight, s.total_closed, s.total_open, s.score_open]
        for c_idx, v in enumerate(vals, 1):
            cell = ws1.cell(row=r_idx, column=c_idx, value=v)
            cell.font = data_font
            cell.border = thin_border
        r_idx += 1

    # Totals
    tot_closed = sum(s.total_closed for s in scores)
    tot_open = sum(s.total_open for s in scores)
    tot_score = sum(s.score_open for s in scores)
    ws1.cell(row=r_idx, column=1, value="Total").font = bold_font
    ws1.cell(row=r_idx, column=3, value=tot_closed).font = bold_font
    ws1.cell(row=r_idx, column=4, value=tot_open).font = bold_font
    ws1.cell(row=r_idx, column=5, value=tot_score).font = bold_font

    # Sheet 2: Defect Details
    ws2 = wb.create_sheet("Defect Details")
    headers_s2 = ["Defect ID", "Module", "Sub-Module", "Summary", "Level", "Priority", "Status", "Aging", "Created By", "Date Created"]
    for c_idx, h in enumerate(headers_s2, 1):
        cell = ws2.cell(row=1, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    defects = db.query(Defect).order_by(Defect.id.desc()).all()
    for row_idx, d in enumerate(defects, 2):
        mod = db.query(Module).filter(Module.id == d.module_id).first()
        sub = db.query(SubModule).filter(SubModule.id == d.sub_module_id).first()
        row_vals = [
            d.defect_id, mod.name if mod else "", sub.name if sub else "",
            d.summary, d.level_of_defect, d.priority, d.status, d.aging,
            d.created_by, str(d.date_created) if d.date_created else ""
        ]
        for c_idx, v in enumerate(row_vals, 1):
            cell = ws2.cell(row=row_idx, column=c_idx, value=v)
            cell.font = data_font
            cell.border = thin_border

    # Sheet 3: Test Scripts
    ws3 = wb.create_sheet("Test Scripts")
    headers_s3 = ["Test Case ID", "Summary", "Module", "Stage", "Tester", "Status", "Completion Date"]
    for c_idx, h in enumerate(headers_s3, 1):
        cell = ws3.cell(row=1, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    scripts = db.query(TestScript).order_by(TestScript.id.desc()).all()
    for row_idx, ts in enumerate(scripts, 2):
        mod = db.query(Module).filter(Module.id == ts.module_id).first()
        row_vals = [
            ts.test_case_id, ts.summary, mod.name if mod else "",
            ts.stage, ts.tester, ts.status_by_tester,
            str(ts.completion_date) if ts.completion_date else ""
        ]
        for c_idx, v in enumerate(row_vals, 1):
            cell = ws3.cell(row=row_idx, column=c_idx, value=v)
            cell.font = data_font
            cell.border = thin_border

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=Executive_SIT_Report_{date.today().strftime('%Y%m%d')}.xlsx"}
    )
