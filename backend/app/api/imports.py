from fastapi import APIRouter, Depends, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, case
from app.database import get_db
from app.models.models import Defect, Module, SubModule, Tester, ReportSummary, DefectScoring
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from io import BytesIO
from datetime import datetime, date
from typing import Optional
import re

router = APIRouter()

HEADER_ROW = [
    "Defect ID#", "Module", "Sub-Module", "Summary", "Stage", "Environment",
    "Description", "Issue Link", "Impact of Issue", "Level of Defect",
    "Scoring Level of Defect", "Priority of Defect Fixing",
    "Defect or Non-Defect Criteria", "Status / Workflow",
    "Date Created", "Date Re-Opened", "Date Closed", "Aging",
    "Created by", "Last Retested by", "Fixing / Confirmed by",
    "Estimated / Actual Fix Date", "Fixing / Review Status",
    "Keterangan", "Retesting"
]

# Known header keywords that mark non-data rows
HEADER_KEYWORDS = {
    "defect id#", "module", "sub-module", "summary", "stage",
    "no", "date", "created*", "re-opened*", "closed*",
    "test case id*", "*",
}

LEVEL_WEIGHTS = {"Fatal": 25, "Major": 10, "Minor": 2, "Kosmetik": 1}
LEVEL_PRIORITIES = {"Fatal": "Highest", "Major": "High", "Minor": "Medium", "Kosmetik": "Low"}
VALID_STATUSES = {"Open", "Closed", "Under Review", "Re-Opened", "Confirmed"}
VALID_LEVELS = {"Fatal", "Major", "Minor", "Kosmetik"}
NON_DEFECT_LABELS = {"Function Running Well", "Application Standards", "User Preference", "Change Request"}


def is_formula(value) -> bool:
    """Check if a cell value is an unresolved Excel formula."""
    if value is None:
        return False
    s = str(value).strip()
    return s.startswith("=") or s.startswith("+") and "(" in s


def clean_cell(value) -> Optional[str]:
    """Clean cell value: strip, handle formula residue, return None for empty."""
    if value is None:
        return None
    s = str(value).strip()
    if s == "" or s == "None" or s == "-" or is_formula(s):
        return None
    return s


def parse_date(value) -> Optional[date]:
    """Parse various date formats from Excel."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    s = str(value).strip()
    if not s or s == "None" or s == "-" or is_formula(s):
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y",
                "%d/%m/%Y %H:%M:%S", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def parse_int(value) -> Optional[int]:
    """Parse integer from Excel cell."""
    if value is None:
        return None
    s = str(value).strip()
    if not s or is_formula(s):
        return None
    try:
        return int(float(s))
    except (ValueError, TypeError):
        return None


def is_header_row(row_values: list) -> bool:
    """Detect if this row is a header/sub-header, not actual data."""
    # Check column A (Defect ID)
    val_a = str(row_values[0] or "").strip().lower()
    # Check column D (Summary)
    val_d = str(row_values[3] or "").strip().lower()

    # Explicit header matches
    if val_a in HEADER_KEYWORDS or val_d in HEADER_KEYWORDS:
        return True
    if val_a in ("defect id#", "no", ""):
        if val_d in ("summary", "defect", "non-defect", "total", "*", ""):
            return True

    # Row where col A is "Defect ID#" or col D is "Summary"
    if "defect id" in val_a or val_d == "summary":
        return True

    # Sub-header row pattern: mostly None/empty or star markers
    non_empty = sum(1 for v in row_values if v is not None and str(v).strip() not in ("", "*"))
    if non_empty <= 3:
        # Likely a spacer or sub-header
        if val_d in ("", "*") and not row_values[1]:
            return True

    return False


def find_data_start_row(ws) -> int:
    """
    Auto-detect where actual defect data starts.
    Scans from row 1 and finds the first row that looks like defect data.
    Returns the 1-based row number.
    """
    for row_idx in range(1, min(ws.max_row + 1, 20)):
        row_values = [ws.cell(row=row_idx, column=c).value for c in range(1, 26)]

        # Skip completely empty rows
        if all(v is None for v in row_values):
            continue

        # Skip header rows
        if is_header_row(row_values):
            continue

        # Check if this looks like a defect data row:
        # - Column D (Summary) has text content (not formula, not header keyword)
        summary = row_values[3]
        if summary is None:
            continue
        summary_str = str(summary).strip()
        if not summary_str or summary_str.lower() in HEADER_KEYWORDS:
            continue

        # - Column B (Module) or C (Sub-Module) has content
        module = row_values[1]
        if module is None or str(module).strip().lower() in HEADER_KEYWORDS:
            continue

        # This row looks like real data
        return row_idx

    # Fallback: try common start rows
    # Row 2 (simple template), Row 3, Row 9 (complex Excel with headers)
    for try_row in [2, 3, 9]:
        if try_row <= ws.max_row:
            val_d = ws.cell(row=try_row, column=4).value
            if val_d and not is_formula(val_d) and str(val_d).strip().lower() not in HEADER_KEYWORDS:
                return try_row

    return 2  # ultimate fallback


def get_or_create_module(db: Session, module_name: str) -> tuple:
    """Get module by name or create it. Returns (id, is_new)."""
    module_name = module_name.strip()
    mod = db.query(Module).filter(Module.name == module_name).first()
    if mod:
        return mod.id, False
    new_mod = Module(name=module_name)
    db.add(new_mod)
    db.flush()
    return new_mod.id, True


def get_or_create_sub_module(db: Session, module_id: int, sub_module_name: str) -> tuple:
    """Get sub-module by name+module or create it. Returns (id, is_new)."""
    sub_module_name = sub_module_name.strip()
    sub = db.query(SubModule).filter(
        SubModule.module_id == module_id,
        SubModule.name == sub_module_name
    ).first()
    if sub:
        return sub.id, False
    new_sub = SubModule(module_id=module_id, name=sub_module_name)
    db.add(new_sub)
    db.flush()
    return new_sub.id, True


def get_or_create_tester(db: Session, tester_name: str) -> int:
    """Get tester by name or create it."""
    tester_name = tester_name.strip()
    tester = db.query(Tester).filter(Tester.name == tester_name).first()
    if tester:
        return tester.id
    new_tester = Tester(name=tester_name)
    db.add(new_tester)
    db.flush()
    return new_tester.id


def sync_report_summary(db: Session):
    """
    Recalculate report_summary from defects table.
    Wipe and rebuild so it always reflects current state.
    """
    db.query(ReportSummary).delete()
    db.flush()

    sub_modules = db.query(SubModule).all()
    for sm in sub_modules:
        defect_count = db.query(func.count(Defect.id)).filter(
            Defect.sub_module_id == sm.id,
            Defect.defect_criteria == "Defect"
        ).scalar() or 0

        non_defect_count = db.query(func.count(Defect.id)).filter(
            Defect.sub_module_id == sm.id,
            Defect.defect_criteria == "Non-Defect"
        ).scalar() or 0

        # Count non-defect sub-categories from level_of_defect field
        frw = db.query(func.count(Defect.id)).filter(
            Defect.sub_module_id == sm.id,
            Defect.defect_criteria == "Non-Defect",
            Defect.level_of_defect == "Function Running Well"
        ).scalar() or 0

        app_std = db.query(func.count(Defect.id)).filter(
            Defect.sub_module_id == sm.id,
            Defect.defect_criteria == "Non-Defect",
            Defect.level_of_defect == "Application Standards"
        ).scalar() or 0

        user_pref = db.query(func.count(Defect.id)).filter(
            Defect.sub_module_id == sm.id,
            Defect.defect_criteria == "Non-Defect",
            Defect.level_of_defect == "User Preference"
        ).scalar() or 0

        change_req = db.query(func.count(Defect.id)).filter(
            Defect.sub_module_id == sm.id,
            Defect.defect_criteria == "Non-Defect",
            Defect.level_of_defect == "Change Request"
        ).scalar() or 0

        rs = ReportSummary(
            sub_module_id=sm.id,
            total_defect=defect_count,
            total_non_defect=non_defect_count,
            total=defect_count + non_defect_count,
            non_defect_function_running_well=frw,
            non_defect_application_standards=app_std,
            non_defect_user_preference=user_pref,
            non_defect_change_request=change_req,
            non_defect_total=frw + app_std + user_pref + change_req,
        )
        db.add(rs)

    db.flush()


def sync_defect_scoring(db: Session):
    """
    Recalculate defect_scoring from defects table.
    """
    db.query(DefectScoring).delete()
    db.flush()

    for category, weight in LEVEL_WEIGHTS.items():
        total_closed = db.query(func.count(Defect.id)).filter(
            Defect.defect_criteria == "Defect",
            Defect.level_of_defect == category,
            Defect.status == "Closed"
        ).scalar() or 0

        total_open = db.query(func.count(Defect.id)).filter(
            Defect.defect_criteria == "Defect",
            Defect.level_of_defect == category,
            Defect.status.in_(["Open", "Re-Opened", "Under Review"])
        ).scalar() or 0

        # Today's defects (created today)
        today = date.today()
        total_today = db.query(func.count(Defect.id)).filter(
            Defect.defect_criteria == "Defect",
            Defect.level_of_defect == category,
            Defect.date_created == today
        ).scalar() or 0

        ds = DefectScoring(
            category=category,
            weight=weight,
            total_defect_today=total_today,
            score_today=total_today * weight,
            total_closed=total_closed,
            total_open=total_open,
            score_open=total_open * weight,
        )
        db.add(ds)

    db.flush()


def sync_testers(db: Session):
    """
    Ensure all tester names from defects exist in testers table.
    """
    # Collect unique names from created_by and last_retested_by
    creators = db.query(Defect.created_by).filter(
        Defect.created_by.isnot(None)
    ).distinct().all()
    retesters = db.query(Defect.last_retested_by).filter(
        Defect.last_retested_by.isnot(None)
    ).distinct().all()

    names = set()
    for (name,) in creators:
        if name and name.strip():
            names.add(name.strip())
    for (name,) in retesters:
        if name and name.strip():
            names.add(name.strip())

    for name in names:
        get_or_create_tester(db, name)

    db.flush()


@router.get("/template")
def download_template():
    """Download Excel template for import."""
    wb = Workbook()
    ws = wb.active
    ws.title = "SIT Import"

    # Styles
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )
    required_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    optional_fill = PatternFill(start_color="475569", end_color="475569", fill_type="solid")
    required_cols = {2, 3, 4, 10, 14, 15, 19}  # B,C,D,J,N,O,S (1-indexed)

    # Write headers at row 1
    for col_idx, header in enumerate(HEADER_ROW, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = required_fill if col_idx in required_cols else optional_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Sample data at row 2
    sample = [
        "", "Vendor Management",
        "Vendor Registration - Disclaimer & Self Registration",
        "Seluruh Template dokumen tidak dapat di akses",
        "Testing: SIT", "Development",
        "Seluruh Template dokumen tidak dapat di akses, ketika di unduh file tidak bisa dibuka",
        "VM-REG_10", "", "Minor", "", "", "Defect", "Closed",
        "2026-07-31", "", "2026-08-04", "",
        "Amanda", "Destra", "Vendor IT", "", "Done",
        "[31/07/2026] Amanda: Defect Created.", ""
    ]
    data_font = Font(name="Calibri", size=10)
    for col_idx, val in enumerate(sample, 1):
        cell = ws.cell(row=2, column=col_idx, value=val)
        cell.font = data_font
        cell.border = thin_border

    # Info sheet
    ws_info = wb.create_sheet("Info")
    info_data = [
        ["Panduan Import Excel - Dashboard SIT"],
        [],
        ["PENTING:"],
        ["- Sistem auto-detect baris data pertama (skip header/rumus di row 1-8)"],
        ["- Defect ID auto-generate jika kosong"],
        ["- Scoring & Priority auto-calculate dari Level of Defect"],
        ["- Aging auto-calculate jika kosong"],
        ["- Jika Defect ID sudah ada, data akan di-UPDATE"],
        ["- Import akan meng-update SEMUA tabel relasi (report, scoring, testers)"],
        [],
        ["Kolom Wajib:"],
        ["  B - Module       : Nama module (dibuat otomatis jika baru)"],
        ["  C - Sub-Module   : Nama sub-module (dibuat otomatis jika baru)"],
        ["  D - Summary      : Ringkasan defect"],
        ["  J - Level        : Fatal / Major / Minor / Kosmetik"],
        ["  N - Status       : Open / Closed / Under Review / Re-Opened / Confirmed"],
        ["  O - Date Created : YYYY-MM-DD atau DD/MM/YYYY"],
        ["  S - Created by   : Nama tester"],
        [],
        ["Level -> Score -> Priority:"],
        ["  Fatal    -> 25 -> Highest"],
        ["  Major    -> 10 -> High"],
        ["  Minor    -> 2  -> Medium"],
        ["  Kosmetik -> 1  -> Low"],
        [],
        ["Non-Defect Criteria (kolom J jika M='Non-Defect'):"],
        ["  Function Running Well, Application Standards,"],
        ["  User Preference, Change Request"],
    ]
    for row_idx, row_data in enumerate(info_data, 1):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws_info.cell(row=row_idx, column=col_idx, value=val)
            if row_idx == 1:
                cell.font = Font(bold=True, size=14)
            elif row_idx == 3:
                cell.font = Font(bold=True, size=11, color="CC0000")

    widths = [12, 25, 45, 50, 15, 15, 50, 15, 20, 15, 12, 12, 15, 15, 15, 15, 15, 8, 15, 15, 15, 18, 18, 45, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws_info.column_dimensions["A"].width = 65

    ws.freeze_panes = "A2"

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=template_import_sit.xlsx"}
    )


@router.post("/excel")
async def import_excel(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """
    Import defects from Excel file.
    - Auto-detects data start row (skips header rows 1-8 or wherever headers end)
    - Handles formula cells that weren't resolved
    - Syncs ALL related tables: report_summary, defect_scoring, testers
    """
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="File harus berformat .xlsx atau .xls")

    try:
        contents = await file.read()

        # Load with data_only=True to get computed values instead of formulas
        wb = load_workbook(BytesIO(contents), data_only=True)
        ws = wb.active

        if ws.max_row < 2:
            raise HTTPException(status_code=400, detail="File kosong atau hanya berisi header")

        # Auto-detect where data starts
        data_start = find_data_start_row(ws)

        results = {
            "total_rows": 0,
            "imported": 0,
            "skipped": 0,
            "updated": 0,
            "errors": [],
            "new_modules": [],
            "new_sub_modules": [],
            "new_testers": [],
            "data_start_row": data_start,
            "synced_tables": [],
        }

        # Get current max defect number for auto-ID generation
        max_num = 0
        existing_ids = db.query(Defect.defect_id).all()
        for (did,) in existing_ids:
            match = re.search(r"(\d+)$", did)
            if match:
                max_num = max(max_num, int(match.group(1)))

        for row_idx in range(data_start, ws.max_row + 1):
            # Read row values (columns A-Y, 1-25)
            row_values = [ws.cell(row=row_idx, column=c).value for c in range(1, 26)]

            # Skip if this looks like a header row that slipped through
            if is_header_row(row_values):
                continue

            # Skip empty rows: require Summary (D, index 3) to have content
            summary_raw = row_values[3]
            if summary_raw is None:
                continue
            summary = clean_cell(summary_raw)
            if not summary:
                continue

            results["total_rows"] += 1
            row_num = row_idx

            try:
                # --- Module & Sub-Module ---
                module_name = clean_cell(row_values[1])
                sub_module_name = clean_cell(row_values[2])

                if not module_name:
                    results["errors"].append(f"Row {row_num}: Module kosong")
                    results["skipped"] += 1
                    continue
                if not sub_module_name:
                    results["errors"].append(f"Row {row_num}: Sub-Module kosong")
                    results["skipped"] += 1
                    continue

                module_id, mod_is_new = get_or_create_module(db, module_name)
                if mod_is_new and module_name not in results["new_modules"]:
                    results["new_modules"].append(module_name)

                sub_module_id, sub_is_new = get_or_create_sub_module(db, module_id, sub_module_name)
                if sub_is_new and sub_module_name not in results["new_sub_modules"]:
                    results["new_sub_modules"].append(sub_module_name)

                # --- Defect ID ---
                defect_id = clean_cell(row_values[0])
                if not defect_id:
                    max_num += 1
                    defect_id = f"Defect-{max_num}"

                # --- Level of Defect ---
                level = clean_cell(row_values[9])
                # Handle non-defect category labels stored in level column
                if level and level in NON_DEFECT_LABELS:
                    # This is a non-defect, level is actually the category
                    pass  # keep level as-is for non-defect categorization

                # --- Scoring & Priority (auto-calculate) ---
                scoring = parse_int(row_values[10])
                priority = clean_cell(row_values[11])

                if level and level in LEVEL_WEIGHTS:
                    if scoring is None:
                        scoring = LEVEL_WEIGHTS[level]
                    if not priority:
                        priority = LEVEL_PRIORITIES[level]

                # --- Criteria ---
                criteria = clean_cell(row_values[12])
                if not criteria or criteria not in ("Defect", "Non-Defect"):
                    if level and level in LEVEL_WEIGHTS:
                        criteria = "Defect"
                    elif level and level in NON_DEFECT_LABELS:
                        criteria = "Non-Defect"
                    else:
                        criteria = "Defect"

                # --- Status ---
                status = clean_cell(row_values[13])
                if not status or status not in VALID_STATUSES:
                    status = "Open"

                # --- Stage & Environment ---
                stage = clean_cell(row_values[4]) or "Testing: SIT"
                environment = clean_cell(row_values[5]) or "Development"

                # --- Dates ---
                date_created = parse_date(row_values[14])
                date_reopened = parse_date(row_values[15])
                date_closed = parse_date(row_values[16])

                # --- Aging (auto-calculate if missing) ---
                aging = parse_int(row_values[17])
                if aging is None and date_created:
                    if date_closed:
                        aging = (date_closed - date_created).days
                    else:
                        aging = (date.today() - date_created).days

                # --- People ---
                created_by = clean_cell(row_values[18])
                last_retested_by = clean_cell(row_values[19])
                fixing_confirmed_by = clean_cell(row_values[20])

                # Track new testers
                for tname in [created_by, last_retested_by]:
                    if tname and tname not in results["new_testers"]:
                        existing_tester = db.query(Tester).filter(Tester.name == tname).first()
                        if not existing_tester:
                            results["new_testers"].append(tname)

                # --- Other fields ---
                description = clean_cell(row_values[6])
                issue_link = clean_cell(row_values[7])
                impact_of_issue = clean_cell(row_values[8])
                estimated_fix_date = parse_date(row_values[21])
                fixing_review_status = clean_cell(row_values[22])
                keterangan = clean_cell(row_values[23])
                retesting = clean_cell(row_values[24])

                # --- Upsert ---
                existing = db.query(Defect).filter(Defect.defect_id == defect_id).first()

                defect_data = dict(
                    module_id=module_id,
                    sub_module_id=sub_module_id,
                    summary=summary,
                    stage=stage,
                    environment=environment,
                    description=description,
                    issue_link=issue_link,
                    impact_of_issue=impact_of_issue,
                    level_of_defect=level,
                    scoring_level=scoring,
                    priority=priority,
                    defect_criteria=criteria,
                    status=status,
                    date_created=date_created,
                    date_reopened=date_reopened,
                    date_closed=date_closed,
                    aging=aging or 0,
                    created_by=created_by,
                    last_retested_by=last_retested_by,
                    fixing_confirmed_by=fixing_confirmed_by,
                    estimated_fix_date=estimated_fix_date,
                    fixing_review_status=fixing_review_status,
                    keterangan=keterangan,
                    retesting=retesting,
                )

                if existing:
                    for key, value in defect_data.items():
                        setattr(existing, key, value)
                    results["updated"] += 1
                else:
                    new_defect = Defect(defect_id=defect_id, **defect_data)
                    db.add(new_defect)
                    results["imported"] += 1

            except Exception as e:
                results["errors"].append(f"Row {row_num}: {str(e)}")
                results["skipped"] += 1
                continue

        # Commit defect changes first
        db.commit()

        # --- Sync all related tables ---
        try:
            sync_testers(db)
            results["synced_tables"].append("testers")
        except Exception as e:
            results["errors"].append(f"Sync testers error: {str(e)}")

        try:
            sync_report_summary(db)
            results["synced_tables"].append("report_summary")
        except Exception as e:
            results["errors"].append(f"Sync report_summary error: {str(e)}")

        try:
            sync_defect_scoring(db)
            results["synced_tables"].append("defect_scoring")
        except Exception as e:
            results["errors"].append(f"Sync defect_scoring error: {str(e)}")

        db.commit()

        return {
            "success": True,
            "message": (
                f"Import selesai: {results['imported']} baru, {results['updated']} diperbarui, "
                f"{results['skipped']} dilewati. "
                f"Data mulai dari row {data_start}. "
                f"Tabel ter-sync: {', '.join(results['synced_tables'])}."
            ),
            **results
        }

    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Gagal memproses file: {str(e)}")


@router.get("/export")
def export_defects(db: Session = Depends(get_db)):
    """Export all defects to Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "SIT Defects"

    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    for col_idx, header in enumerate(HEADER_ROW, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    defects = db.query(Defect).order_by(Defect.id).all()
    data_font = Font(name="Calibri", size=10)

    for row_idx, d in enumerate(defects, 2):
        module = db.query(Module).filter(Module.id == d.module_id).first()
        sub_module = db.query(SubModule).filter(SubModule.id == d.sub_module_id).first()

        values = [
            d.defect_id,
            module.name if module else "",
            sub_module.name if sub_module else "",
            d.summary, d.stage, d.environment, d.description,
            d.issue_link, d.impact_of_issue, d.level_of_defect,
            d.scoring_level, d.priority, d.defect_criteria, d.status,
            str(d.date_created) if d.date_created else "",
            str(d.date_reopened) if d.date_reopened else "",
            str(d.date_closed) if d.date_closed else "",
            d.aging, d.created_by, d.last_retested_by,
            d.fixing_confirmed_by,
            str(d.estimated_fix_date) if d.estimated_fix_date else "",
            d.fixing_review_status, d.keterangan, d.retesting,
        ]

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border

    widths = [12, 25, 45, 50, 15, 15, 50, 15, 20, 15, 12, 12, 15, 15, 15, 15, 15, 8, 15, 15, 15, 18, 18, 45, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    ws.freeze_panes = "A2"

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=export_sit_{date.today().strftime('%Y%m%d')}.xlsx"}
    )
