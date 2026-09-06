from datetime import date
from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.models import Defect, Module, SubModule, TestCase

router = APIRouter()

REPORT_HEADERS = [
    "Defect ID", "Nama File Import", "Modul", "Sheet Name", "Summary",
    "Priority of Defect Fixing", "Status Workflow", "Date Created",
    "Fixing / Review Status", "Fixing Status by Vendor", "Keterangan", "Note",
]


def base_query(db: Session):
    return db.query(Defect, Module.name.label("module_name"), SubModule.name.label("sub_module_name")).outerjoin(
        Module, Module.id == Defect.module_id
    ).outerjoin(SubModule, SubModule.id == Defect.sub_module_id).filter(Defect.status == "Open")


def apply_filters(query, module_id: Optional[int], search: Optional[str],
                  fixing_status: Optional[str], vendor_status: Optional[str],
                  date_from: Optional[date], date_to: Optional[date]):
    if module_id:
        query = query.filter(Defect.sub_module_id == module_id)
    if fixing_status:
        query = query.filter(Defect.fixing_review_status == fixing_status)
    if vendor_status:
        query = query.filter(Defect.fixing_status_by_vendor == vendor_status)
    if date_from:
        query = query.filter(Defect.date_created >= date_from)
    if date_to:
        query = query.filter(Defect.date_created <= date_to)
    if search:
        term = f"%{search}%"
        query = query.filter(Defect.defect_id.ilike(term) | Defect.summary.ilike(term))
    return query


def test_case_metadata(db: Session, defects):
    """Resolve test metadata from the existing issue and sub-module relations."""
    ids = set()
    sub_module_ids = {defect.sub_module_id for defect in defects if defect.sub_module_id}
    for defect in defects:
        ids.update(
            token.strip().lower()
            for token in (defect.issue_link or "").split(",")
            if token.strip()
        )
    cases = db.query(TestCase, Module.name.label("module_name")).outerjoin(
        Module, Module.id == TestCase.module_id
    ).filter(TestCase.test_case_id.isnot(None)).all()
    metadata = {"by_issue_link": {}, "by_sub_module": {}}
    sub_modules = {
        sub_module.id: sub_module.name.strip().lower()
        for sub_module in db.query(SubModule).filter(SubModule.id.in_(sub_module_ids)).all()
    }
    for case, module_name in cases:
        case_id = case.test_case_id.strip().lower()
        if case_id in ids and case_id not in metadata["by_issue_link"]:
            metadata["by_issue_link"][case_id] = (case.import_file_name, case.sheet_name)
        module_key = (module_name or "").strip().lower()
        for sub_module_id, sub_module_name in sub_modules.items():
            if module_key == sub_module_name and sub_module_id not in metadata["by_sub_module"]:
                metadata["by_sub_module"][sub_module_id] = (case.import_file_name, case.sheet_name)
    return metadata


def report_row(defect, module_name, sub_module_name, test_metadata):
    file_name = sheet_name = None
    for token in (defect.issue_link or "").split(","):
        match = test_metadata["by_issue_link"].get(token.strip().lower())
        if match:
            file_name = file_name or match[0]
            sheet_name = sheet_name or match[1]
    if not file_name and not sheet_name:
        match = test_metadata["by_sub_module"].get(defect.sub_module_id)
        if match:
            file_name, sheet_name = match
    return {
        "defect_id": defect.defect_id,
        "import_file_name": file_name,
        "module_name": sub_module_name or module_name,
        "sheet_name": sheet_name,
        "summary": defect.summary,
        "priority": defect.priority,
        "status": defect.status,
        "date_created": defect.date_created,
        "fixing_review_status": defect.fixing_review_status,
        "fixing_status_by_vendor": defect.fixing_status_by_vendor,
        "keterangan": defect.keterangan,
        "note": defect.retesting,
    }


@router.get("/")
def get_defect_intake(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    module_id: Optional[int] = None,
    search: Optional[str] = None,
    fixing_status: Optional[str] = None,
    vendor_status: Optional[str] = None,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    import_file_name: Optional[str] = None,
    db: Session = Depends(get_db),
):
    query = apply_filters(base_query(db), module_id, search, fixing_status,
                          vendor_status, date_from, date_to)
    records = query.order_by(Defect.date_created.desc(), Defect.id.desc()).all()
    metadata = test_case_metadata(db, [defect for defect, _, _ in records])
    if import_file_name:
        records = [record for record in records if report_row(record[0], record[1], record[2], metadata)["import_file_name"] == import_file_name]
    total = len(records)
    records = records[(page - 1) * page_size:page * page_size]
    return {
        "items": [report_row(defect, module_name, sub_module_name, metadata) for defect, module_name, sub_module_name in records],
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_pages": max(1, (total + page_size - 1) // page_size),
    }


@router.get("/filter-options")
def get_defect_intake_filter_options(db: Session = Depends(get_db)):
    rows = base_query(db).all()
    metadata = test_case_metadata(db, [defect for defect, _, _ in rows])
    reports = [report_row(defect, module_name, sub_module_name, metadata) for defect, module_name, sub_module_name in rows]
    return {
        "modules": [{"id": module.id, "name": module.name} for module in db.query(SubModule).order_by(SubModule.name).all()],
        "fixing_statuses": sorted({defect.fixing_review_status for defect, _, _ in rows if defect.fixing_review_status}),
        "vendor_statuses": sorted({defect.fixing_status_by_vendor for defect, _, _ in rows if defect.fixing_status_by_vendor}),
        "import_files": sorted({row["import_file_name"] for row in reports if row["import_file_name"]}),
    }


@router.get("/summary")
def get_defect_intake_summary(
    module_id: Optional[int] = None,
    search: Optional[str] = None,
    fixing_status: Optional[str] = None,
    vendor_status: Optional[str] = None,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    import_file_name: Optional[str] = None,
    db: Session = Depends(get_db),
):
    """Return live Open-defect distributions using the same filters as the table."""
    query = apply_filters(base_query(db), module_id, search, fixing_status,
                          vendor_status, date_from, date_to)
    rows = query.all()
    metadata = test_case_metadata(db, [defect for defect, _, _ in rows])
    filtered = [report_row(defect, module_name, sub_module_name, metadata) for defect, module_name, sub_module_name in rows]
    if import_file_name:
        filtered = [row for row in filtered if row["import_file_name"] == import_file_name]

    def distribution(index, fallback):
        counts = {}
        for row in filtered:
            label = row[("priority", "fixing_review_status", "fixing_status_by_vendor")[index]] or fallback
            counts[label] = counts.get(label, 0) + 1
        return [{"label": label, "count": count} for label, count in sorted(counts.items(), key=lambda item: (-item[1], item[0]))]

    return {
        "total": len(filtered),
        "priority": distribution(0, "Belum diisi"),
        "fixing_review": distribution(1, "Belum diisi"),
        "vendor": distribution(2, "Belum diisi"),
    }


@router.get("/export")
def export_defect_intake(
    module_id: Optional[int] = None,
    search: Optional[str] = None,
    fixing_status: Optional[str] = None,
    vendor_status: Optional[str] = None,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    import_file_name: Optional[str] = None,
    db: Session = Depends(get_db),
):
    query = apply_filters(base_query(db), module_id, search, fixing_status,
                          vendor_status, date_from, date_to)
    records = query.order_by(Defect.date_created.desc(), Defect.id.desc()).all()
    metadata = test_case_metadata(db, [defect for defect, _, _ in records])
    if import_file_name:
        records = [record for record in records if report_row(record[0], record[1], record[2], metadata)["import_file_name"] == import_file_name]

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Defect Intake"
    header_fill = PatternFill("solid", fgColor="334155")
    border = Border(*(Side(style="thin", color="CBD5E1"),) * 4)
    soft_fills = {
        "Highest": PatternFill("solid", fgColor="FCE7F3"),
        "High": PatternFill("solid", fgColor="DBEAFE"),
        "Medium": PatternFill("solid", fgColor="FEF3C7"),
        "Low": PatternFill("solid", fgColor="E2E8F0"),
        "Fix in Progress": PatternFill("solid", fgColor="DBEAFE"),
        "Ready to Test": PatternFill("solid", fgColor="EDE9FE"),
        "Needs Attention": PatternFill("solid", fgColor="FEE2E2"),
        "Critical Action Needed": PatternFill("solid", fgColor="FCE7F3"),
        "Done Dev": PatternFill("solid", fgColor="D1FAE5"),
        "Belum diisi": PatternFill("solid", fgColor="F1F5F9"),
    }
    for column, header in enumerate(REPORT_HEADERS, 1):
        cell = worksheet.cell(1, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row_number, (defect, module_name, sub_module_name) in enumerate(records, 2):
        data = report_row(defect, module_name, sub_module_name, metadata)
        values = [
            data["defect_id"], data["import_file_name"], data["module_name"], data["sheet_name"],
            data["summary"], data["priority"], data["status"], data["date_created"],
            data["fixing_review_status"], data["fixing_status_by_vendor"], data["keterangan"], data["note"],
        ]
        for column, value in enumerate(values, 1):
            cell = worksheet.cell(row_number, column, value)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if column == 8 and value:
                cell.number_format = "DD/MM/YYYY"
            if column in (6, 7, 9, 10):
                cell.fill = soft_fills.get(value, soft_fills["Belum diisi"])
                cell.font = Font(bold=True, color="475569")
    widths = [16, 26, 30, 22, 55, 24, 18, 16, 25, 25, 45, 45]
    for column, width in enumerate(widths, 1):
        worksheet.column_dimensions[worksheet.cell(1, column).column_letter].width = width
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.freeze_panes = "A2"
    worksheet.row_dimensions[1].height = 32

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=defect_intake_report.xlsx"},
    )
